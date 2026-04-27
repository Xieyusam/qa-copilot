"""
Unit tests for DocumentParser (task 3.1).
Covers: empty files, TXT/MD/DOCX/PDF parsing, error handling.
"""
from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import pytest

from app.core.schemas import ParsedDocument
from app.services.document.parser import DocumentParser


@pytest.fixture
def parser() -> DocumentParser:
    return DocumentParser()


# ---------------------------------------------------------------------------
# Helper: write a temp file and return its path
# ---------------------------------------------------------------------------

def _tmp(suffix: str, content: bytes | str) -> str:
    fd, path = tempfile.mkstemp(suffix=suffix)
    try:
        if isinstance(content, str):
            content = content.encode("utf-8")
        os.write(fd, content)
    finally:
        os.close(fd)
    return path


# ---------------------------------------------------------------------------
# Empty file tests (requirement 5.4)
# ---------------------------------------------------------------------------

class TestEmptyFile:
    def test_empty_txt_returns_empty_parsed_document(self, parser):
        path = _tmp(".txt", "")
        try:
            result = parser.parse(path, "txt")
            assert isinstance(result, ParsedDocument)
            assert result.content == ""
            assert result.paragraphs == []
        finally:
            os.unlink(path)

    def test_empty_md_returns_empty_parsed_document(self, parser):
        path = _tmp(".md", "")
        try:
            result = parser.parse(path, "md")
            assert isinstance(result, ParsedDocument)
            assert result.content == ""
            assert result.paragraphs == []
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# TXT parsing
# ---------------------------------------------------------------------------

class TestTxtParsing:
    def test_single_paragraph(self, parser):
        path = _tmp(".txt", "Hello world")
        try:
            result = parser.parse(path, "txt")
            assert result.content == "Hello world"
            assert result.paragraphs == ["Hello world"]
            assert result.file_type == "txt"
        finally:
            os.unlink(path)

    def test_double_newline_splits_paragraphs(self, parser):
        path = _tmp(".txt", "Para one\n\nPara two\n\nPara three")
        try:
            result = parser.parse(path, "txt")
            assert len(result.paragraphs) == 3
            assert result.paragraphs[0] == "Para one"
            assert result.paragraphs[2] == "Para three"
        finally:
            os.unlink(path)

    def test_single_newline_fallback(self, parser):
        path = _tmp(".txt", "Line one\nLine two")
        try:
            result = parser.parse(path, "txt")
            assert len(result.paragraphs) == 2
        finally:
            os.unlink(path)

    def test_doc_id_propagated(self, parser):
        path = _tmp(".txt", "content")
        try:
            result = parser.parse(path, "txt", doc_id="my-id")
            assert result.doc_id == "my-id"
        finally:
            os.unlink(path)

    def test_doc_id_generated_when_not_provided(self, parser):
        path = _tmp(".txt", "content")
        try:
            result = parser.parse(path, "txt")
            assert result.doc_id  # non-empty UUID
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------

class TestMarkdownParsing:
    def test_md_file_type_stored(self, parser):
        path = _tmp(".md", "# Title\n\nSome text")
        try:
            result = parser.parse(path, "md")
            assert result.file_type == "md"
            assert len(result.paragraphs) >= 1
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# DOCX parsing
# ---------------------------------------------------------------------------

class TestDocxParsing:
    def _make_docx(self, paragraphs: list[str]) -> str:
        from docx import Document
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(path)
        return path

    def test_docx_paragraphs_extracted(self, parser):
        path = self._make_docx(["First paragraph", "Second paragraph"])
        try:
            result = parser.parse(path, "docx")
            assert result.file_type == "docx"
            assert "First paragraph" in result.paragraphs
            assert "Second paragraph" in result.paragraphs
        finally:
            os.unlink(path)

    def test_empty_docx_returns_empty_document(self, parser):
        path = self._make_docx([])
        try:
            result = parser.parse(path, "docx")
            assert result.content == ""
            assert result.paragraphs == []
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# PDF parsing
# ---------------------------------------------------------------------------

class TestPdfParsing:
    def _make_pdf(self, text: str) -> str:
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), text)
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        doc.save(path)
        doc.close()
        return path

    def test_pdf_content_extracted(self, parser):
        path = self._make_pdf("Hello from PDF")
        try:
            result = parser.parse(path, "pdf")
            assert result.file_type == "pdf"
            assert "Hello from PDF" in result.content
            assert result.metadata.get("page_count") == 1
        finally:
            os.unlink(path)

    def test_pdf_paragraphs_non_empty(self, parser):
        path = self._make_pdf("Paragraph one\n\nParagraph two")
        try:
            result = parser.parse(path, "pdf")
            assert len(result.paragraphs) >= 1
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Error handling (requirement 2.4)
# ---------------------------------------------------------------------------

class TestErrorHandling:
    def test_corrupt_pdf_does_not_raise(self, parser):
        path = _tmp(".pdf", b"not a real pdf")
        try:
            result = parser.parse(path, "pdf")
            assert isinstance(result, ParsedDocument)
            assert "error" in result.metadata
        finally:
            os.unlink(path)

    def test_unsupported_type_returns_empty_document(self, parser):
        path = _tmp(".xyz", "some content")
        try:
            result = parser.parse(path, "xyz")
            assert isinstance(result, ParsedDocument)
            assert result.content == ""
            assert "error" in result.metadata
        finally:
            os.unlink(path)

    def test_nonexistent_file_does_not_raise(self, parser):
        result = parser.parse("/nonexistent/path/file.txt", "txt")
        assert isinstance(result, ParsedDocument)
        assert "error" in result.metadata

# ---------------------------------------------------------------------------
# Excel (XLSX) parsing
# ---------------------------------------------------------------------------

class TestExcelParsing:
    def _make_xlsx(self, data: dict) -> str:
        import pandas as pd
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        df = pd.DataFrame(data)
        df.to_excel(path, index=False)
        return path

    def test_xlsx_content_extracted(self, parser):
        data = {"Name": ["Alice", "Bob"], "Age": [25, 30]}
        path = self._make_xlsx(data)
        try:
            result = parser.parse(path, "xlsx")
            assert result.file_type == "xlsx"
            assert "Name: Alice" in result.content
            assert "Age: 25" in result.content
            assert "Name: Bob" in result.content
            assert "Age: 30" in result.content
        finally:
            os.unlink(path)

    def test_empty_xlsx_returns_empty_document(self, parser):
        path = self._make_xlsx({})
        try:
            result = parser.parse(path, "xlsx")
            assert result.content == ""
            assert result.paragraphs == []
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Metadata fields
# ---------------------------------------------------------------------------

class TestMetadata:
    def test_txt_metadata_has_paragraph_count(self, parser):
        path = _tmp(".txt", "Para one\n\nPara two")
        try:
            result = parser.parse(path, "txt")
            assert result.metadata["paragraph_count"] == len(result.paragraphs)
        finally:
            os.unlink(path)
